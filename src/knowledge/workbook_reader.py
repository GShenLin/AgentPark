"""One workbook boundary for index ingestion, source previews and exact aggregation.

XLSX streams rows; XLS loads/unloads one sheet at a time. Never evaluates macros,
external links or formulas. Formula text and saved results are kept separately.
"""
from contextlib import contextmanager
from datetime import date, datetime, time
from itertools import zip_longest
import math

from .document_types import DocumentError
from .office_limits import check_office_zip
from .scanner import check_stop


MAX_CELLS = 5_000_000


def scalar(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise DocumentError("Excel 包含非有限数值")
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    raise DocumentError(f"不支持的 Excel 单元格值类型：{type(value).__name__}")


class Workbook:
    def __init__(self, path, config, stop):
        self.stop = stop
        self.book = self.values = None
        self.xlsx = path.suffix.lower() == ".xlsx"
        try:
            if self.xlsx:
                from openpyxl import load_workbook
                check_office_zip(path, config.max_file_mb)
                self.book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                self.values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            else:
                import xlrd
                self.book = xlrd.open_workbook(str(path), on_demand=True, ragged_rows=True)
        except BaseException:
            self.close()
            raise

    @property
    def sheet_names(self):
        return self.book.sheetnames if self.xlsx else self.book.sheet_names()

    def metadata(self):
        return {"sheets": self.sheet_names,
                "formula_policy": "saved_values_only; formulas are not recalculated",
                "merged_cell_policy": "only the anchor contains a value; no automatic fill",
                "hidden_policy": "includes hidden rows, columns and sheets",
                "format": "xlsx" if self.xlsx else "xls",
                "warnings": ["Excel 使用文件中保存的值；公式结果可能过期。图表、图片、宏不参与解析。"] +
                ([] if self.xlsx else ["XLS 读取保存的计算结果，无法区分普通值和公式，也无法检测公式缓存缺失。"])}

    def rows(self, sheet, start=1, end=None, max_column=None):
        if sheet not in self.sheet_names:
            raise DocumentError(f"工作表不存在：{sheet}")
        check_stop(self.stop)
        iterator = self._xlsx_rows(sheet, start, end, max_column) if self.xlsx else self._xls_rows(sheet, start, end, max_column)
        cells = 0
        try:
            for number, row in iterator:
                check_stop(self.stop)
                cells += len(row)
                if cells > MAX_CELLS:
                    raise DocumentError("单次读取超过 500 万个单元格；请指定更小的行列范围")
                yield number, row
        finally:
            iterator.close()

    def _xlsx_rows(self, sheet, start, end, max_column):
        from openpyxl.utils import get_column_letter
        formulas, values = self.book[sheet], self.values[sheet]
        # Do not trust producers that incorrectly write dimension=A1:A1.
        formulas.reset_dimensions()
        values.reset_dimensions()
        args = dict(min_row=start, max_row=end, max_col=max_column)
        left, right = formulas.iter_rows(**args), values.iter_rows(**args)
        try:
            for number, pair in enumerate(zip_longest(left, right), start):
                if pair[0] is None or pair[1] is None or len(pair[0]) != len(pair[1]):
                    raise DocumentError("Excel 公式与缓存行不一致")
                row = []
                for column, (raw, cached) in enumerate(zip(*pair), 1):
                    formula = raw.value if raw.data_type == "f" else None
                    if formula is not None and not isinstance(formula, str):
                        formula = getattr(formula, "text", None)
                        if not isinstance(formula, str):
                            raise DocumentError(f"{sheet}!{get_column_letter(column)}{number} 的公式类型不受支持")
                    value = scalar(cached.value if formula is not None else raw.value)
                    error = (cached.value if formula is not None else raw.value) if (
                        cached.data_type == "e" if formula is not None else raw.data_type == "e") else None
                    row.append({"column": get_column_letter(column), "value": value,
                                "formula": formula, "error": error,
                                "missing_cache": formula is not None and cached.value is None})
                yield number, row
        finally:
            left.close()
            right.close()

    def _xls_rows(self, sheet, start, end, max_column):
        import xlrd
        from openpyxl.utils import get_column_letter
        index = self.sheet_names.index(sheet)
        tab = self.book.sheet_by_index(index)
        try:
            for row_index in range(start - 1, min(end or tab.nrows, tab.nrows)):
                check_stop(self.stop)
                row = []
                for column in range(min(tab.row_len(row_index), max_column or tab.ncols)):
                    cell = tab.cell(row_index, column)
                    value, error = cell.value, None
                    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, self.book.datemode).isoformat()
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        error = xlrd.error_text_from_code.get(value, f"Excel error {value}")
                        value = error
                    row.append({"column": get_column_letter(column + 1), "value": scalar(value),
                                "formula": None, "error": error, "missing_cache": False})
                yield row_index + 1, row
        finally:
            self.book.unload_sheet(index)

    def close(self):
        for book in (self.book, self.values):
            if book is not None:
                if self.xlsx:
                    book.close()
                else:
                    book.release_resources()


@contextmanager
def open_workbook(path, config, stop):
    book = Workbook(path, config, stop)
    try:
        yield book
    finally:
        book.close()


def spreadsheet_passages(path, config, stop, warnings):
    from .parser import split_text
    from .document_types import Passage
    with open_workbook(path, config, stop) as book:
        warnings.extend(book.metadata()["warnings"])
        missing = errors = 0
        for sheet in book.sheet_names:
            headers = {}
            buffer, location = "", None
            for number, row in book.rows(sheet):
                if number == config.spreadsheet_header_row:
                    headers = {cell["column"]: str(cell["value"]) for cell in row
                               if cell["value"] is not None and not cell["error"] and not cell["missing_cache"]}
                cells = []
                for cell in row:
                    if cell["value"] is None and cell["formula"] is None:
                        continue
                    value = str(cell["value"]) if cell["value"] is not None else "[未保存计算结果]"
                    if cell["formula"] is not None:
                        value += f" (公式: {cell['formula']})"
                    label = f" ({headers[cell['column']]})" if cell["column"] in headers and number != config.spreadsheet_header_row else ""
                    cells.append(f"{cell['column']}{number}{label}: {value}")
                    missing += int(cell["missing_cache"])
                    errors += int(bool(cell["error"]))
                if not cells:
                    continue
                text = f"工作表：{sheet}；第 {number} 行\n" + " | ".join(cells)
                if len(text) > config.max_page_chars:
                    raise DocumentError(f"{sheet} 第 {number} 行超过字符上限")
                if buffer and len(buffer) + len(text) + 1 > config.chunk_size:
                    yield Passage(buffer, location=location)
                    buffer, location = "", None
                for piece in split_text(text, config.chunk_size, config.chunk_overlap):
                    piece.location = {"kind": "spreadsheet", "sheet": sheet, "row_start": number,
                                      "row_end": number, "column_start": row[0]["column"],
                                      "column_end": row[-1]["column"]}
                    if config.spreadsheet_header_row:
                        piece.location["header_row"] = config.spreadsheet_header_row
                    if len(text) > config.chunk_size:
                        yield piece
                    else:
                        buffer += ("\n" if buffer else "") + piece.text
                        if location is None:
                            location = piece.location
                        else:
                            location["row_end"] = number
                            from openpyxl.utils import column_index_from_string
                            if column_index_from_string(row[-1]["column"]) > column_index_from_string(location["column_end"]):
                                location["column_end"] = row[-1]["column"]
            if buffer:
                yield Passage(buffer, location=location)
        if missing:
            warnings.append(f"{missing} 个公式未保存计算结果，仅索引公式文本；统计这些单元格时将报错")
        if errors:
            warnings.append(f"{errors} 个单元格包含 Excel 错误值；统计这些单元格时将报错")
