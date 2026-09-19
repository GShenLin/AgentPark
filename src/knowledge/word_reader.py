"""Word text in document order, with stable paragraph/table citations."""
from contextlib import nullcontext

from .document_types import DocumentError
from .legacy_word import converted_doc
from .office_limits import check_office_zip
from .scanner import check_stop


def word_passages(path, config, stop, warnings):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    context = converted_doc(path, stop) if path.suffix.lower() == ".doc" else nullcontext(path)
    with context as source:
        check_stop(stop)
        check_office_zip(source, config.max_file_mb)
        document = Document(source)
        warnings.append("Word 提取正文、表格、页眉和页脚；图片、图表、文本框及修订记录不作为完整文本解析，不提供渲染页码")
        section = ""
        counters = {"paragraph": 0, "table": 0}

        def blocks(container, part):
            nonlocal section
            for block in container.iter_inner_content():
                check_stop(stop)
                if isinstance(block, Paragraph):
                    counters["paragraph"] += 1
                    if block.style and block.style.name.startswith("Heading"):
                        section = block.text
                    location = {"kind": "word", "part": part, "section": section,
                                "paragraph": counters["paragraph"]}
                    yield from pieces(block.text, location)
                elif isinstance(block, Table):
                    counters["table"] += 1
                    table_id = counters["table"]
                    for row_id, row in enumerate(block.rows, 1):
                        check_stop(stop)
                        # A merged cell's text is emitted once per physical row.
                        seen, cells = set(), []
                        for column, cell in enumerate(row.cells, 1):
                            if cell._tc in seen:
                                continue
                            seen.add(cell._tc)
                            cells.append(f"列 {column}: {cell.text}")
                        yield from pieces(" | ".join(cells), {
                            "kind": "word", "part": part, "section": section,
                            "table": table_id, "row_start": row_id, "row_end": row_id})
                        for cell in row.cells:
                            if cell._tc not in seen:
                                continue
                            seen.remove(cell._tc)
                            for nested in cell.tables:
                                # Nested tables are not included by cell.text.
                                yield from table_nested(nested, part, table_id, row_id)

        def table_nested(table, part, parent, row):
            for index, nested_row in enumerate(table.rows, 1):
                check_stop(stop)
                yield from pieces(" | ".join(c.text for c in nested_row.cells), {
                    "kind": "word", "part": part, "section": section,
                    "table": parent, "row_start": row, "row_end": row, "nested_row": index})
                for cell in nested_row.cells:
                    for child in cell.tables:
                        yield from table_nested(child, part, parent, row)

        def pieces(text, location):
            from .parser import split_text
            if not text.strip():
                return
            if len(text) > config.max_page_chars:
                raise DocumentError("Word 单段或表格行超过字符上限，未发布不完整索引")
            prefix = f"章节：{section}\n" if section and text != section else ""
            for piece in split_text(prefix + text, config.chunk_size, config.chunk_overlap):
                piece.location = location
                yield piece

        yield from blocks(document, "body")
        visited = set()
        for index, segment in enumerate(document.sections, 1):
            for name in ("header", "footer", "first_page_header", "first_page_footer", "even_page_header", "even_page_footer"):
                part = getattr(segment, name)
                if part.is_linked_to_previous or part.part.partname in visited:
                    continue
                visited.add(part.part.partname)
                section = ""
                yield from blocks(part, f"section_{index}_{name}")
