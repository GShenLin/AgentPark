import json
from pathlib import Path
import threading
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from openpyxl import Workbook
import pytest

from src.knowledge.catalog import Catalog
from src.knowledge.contracts import LibraryConfig
from src.knowledge.database import Database
from src.knowledge.parser import parse_pending, passages
from src.knowledge.retrieval import search
from src.knowledge.scanner import begin_scan, scan, prune, Interrupted
from src.knowledge.skill import execute


@pytest.fixture
def office(tmp_path):
    folder = tmp_path / "sources"
    folder.mkdir()
    catalog = Catalog(tmp_path)
    config = LibraryConfig(name="Office", folder=str(folder), embedding_url="http://localhost:1/v1",
                           embedding_model="test", dimensions=8, chunk_size=256, chunk_overlap=32)
    key = catalog.create(config)["id"]
    store = Database(catalog.directory(key))
    store.initialize()
    return tmp_path, folder, key, store, config


def ingest(office):
    _, folder, _, store, config = office
    stop = threading.Event()
    begin_scan(store)
    scan(store, folder, stop)
    prune(store, stop)
    parse_pending(store, config, stop)
    return store.documents()["items"]


def call(office, action, **kwargs):
    workspace, _, key, _, _ = office
    return execute(workspace, {"action": action, "library_id": key, **kwargs})


def make_sheet(path):
    book = Workbook()
    tab = book.active
    tab.title = "销售"
    for row in [("地区", "状态", "金额"), ("华东", "完成", 0.1), ("华东", "完成", 0.2),
                ("华北", "完成", 12), ("华北", "取消", 99), ("华东", "完成", None)]:
        tab.append(row)
    tab.row_dimensions[4].hidden = True
    book.create_sheet("备注").append(["第二张表"])
    book.save(path)
    book.close()


def test_word_body_order_tables_headers_locations_and_source_paging(office):
    _, folder, _, store, config = office
    word = Document()
    word.add_heading("技术原理", 1)
    word.add_paragraph("第一段正文")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "检索算法"
    word.add_paragraph("表后正文")
    word.sections[0].header.paragraphs[0].text = "页眉来源"
    word.save(folder / "技术.docx")
    docs = ingest(office)
    assert docs[0]["state"] == "embedding"
    with store.connect() as db:
        row = db.execute("SELECT id,revision FROM documents").fetchone()
        store.publish(db, row["id"], row["revision"])
    hit = search(store, config, "检索算法", mode="keyword")["matches"][0]
    assert hit["location"]["table"] == 1
    assert hit["location"]["section"] == "技术原理"
    assert hit["page"] is None
    first = call(office, "read", document_id=docs[0]["id"], limit=2)
    assert first["next_offset"] == 2
    rest = call(office, "read", document_id=docs[0]["id"], offset=first["next_offset"])
    assert "检索算法" in rest["passages"][0]["text"]
    assert "表后正文" in rest["passages"][1]["text"]
    assert any("页眉来源" in item["text"] for item in rest["passages"])
    assert first["warnings"]


def test_excel_index_locations_listing_and_paged_source(office):
    _, folder, _, store, _ = office
    make_sheet(folder / "销售.xlsx")
    (folder / "~$销售.xlsx").write_bytes(b"lock file")
    docs = ingest(office)
    assert len(docs) == 1
    result = call(office, "list")
    assert result["items"][0]["id"] == docs[0]["id"]
    meta = call(office, "read", document_id=docs[0]["id"])
    assert meta["sheets"] == ["销售", "备注"]
    page = call(office, "read", document_id=docs[0]["id"], sheet="销售", offset=1, limit=2)
    assert [row["row"] for row in page["rows"]] == [2, 3]
    assert page["next_offset"] == 3
    assert page["rows"][0]["cells"][2]["value"] == 0.1
    with store.connect() as db:
        location = json.loads(db.execute("SELECT location FROM chunks ORDER BY id LIMIT 1").fetchone()[0])
    assert location["kind"] == "spreadsheet" and location["sheet"] == "销售"
    assert location["row_start"] == 1 and location["row_end"] >= 1
    assert location["column_start"] == "A" and location["column_end"] == "C"


def test_full_range_decimal_aggregation_filters_hidden_rows_and_blanks(office):
    make_sheet(office[1] / "销售.xlsx")
    key = ingest(office)[0]["id"]
    result = call(office, "table", document_id=key, sheet="销售", start_row=2,
                  filters=[dict(column="B", op="eq", value="完成")], group_by=["A"],
                  metrics=[dict(name="金额", op="sum", column="C"), dict(name="均值", op="mean", column="C"),
                           dict(name="行数", op="count_rows"), dict(name="最大", op="max", column="C")])
    assert result["complete"] and result["matched_rows"] == 4
    east, north = result["groups"]
    assert east["metrics"]["金额"] == {"value": "0.3", "nonblank_count": 2, "blank_count": 1}
    assert east["metrics"]["均值"]["value"] == "0.15"
    assert east["metrics"]["行数"]["value"] == 3
    assert north["metrics"]["金额"]["value"] == "12"


def test_missing_formula_cache_is_visible_and_fails_calculation(office):
    book = Workbook()
    book.active.append(["amount"])
    book.active.append(["=1+2"])
    book.save(office[1] / "formula.xlsx")
    docs = ingest(office)
    assert "未保存计算结果" in docs[0]["warning"]
    read = call(office, "read", document_id=docs[0]["id"], sheet="Sheet", offset=1)
    assert read["rows"][0]["cells"][0]["missing_cache"] is True
    with pytest.raises(ValueError, match="Sheet!A2.*公式缺少"):
        call(office, "table", document_id=docs[0]["id"], sheet="Sheet", start_row=2,
             metrics=[dict(name="sum", op="sum", column="A")])


def test_cached_formula_value_and_formula_are_both_preserved(office):
    path = office[1] / "cached.xlsx"
    book = Workbook()
    book.active.append(["=1+2"])
    book.save(path)
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(b"<v></v>", b"<v>3</v>")
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    key = ingest(office)[0]["id"]
    row = call(office, "read", document_id=key, sheet="Sheet")["rows"][0]
    assert row["cells"][0]["formula"] == "=1+2"
    assert row["cells"][0]["value"] == 3
    result = call(office, "table", document_id=key, sheet="Sheet", metrics=[dict(name="sum", op="sum", column="A")])
    assert result["groups"][0]["metrics"]["sum"]["value"] == "3"


def test_bad_numeric_values_fail_instead_of_being_silently_dropped(office):
    book = Workbook()
    book.active.append(["123"])
    book.save(office[1] / "text.xlsx")
    key = ingest(office)[0]["id"]
    with pytest.raises(ValueError, match="不是数值"):
        call(office, "table", document_id=key, sheet="Sheet", metrics=[dict(name="sum", op="sum", column="A")])


def test_source_updates_and_library_permissions_apply_to_read_and_table(office):
    from src.knowledge.contracts import LibraryUpdate
    workspace, folder, library_id, _, config = office
    path = folder / "销售.xlsx"
    make_sheet(path)
    key = ingest(office)[0]["id"]
    path.write_bytes(path.read_bytes() + b"modified")
    for action, args in [("read", {}), ("table", dict(sheet="销售", metrics=[dict(name="rows", op="count_rows")]))]:
        with pytest.raises(ValueError, match="源文件已变化"):
            call(office, action, document_id=key, **args)
    Catalog(workspace).update(library_id, LibraryUpdate(**{**config.model_dump(), "allow_agents": False}))
    for action in ["read", "list"]:
        with pytest.raises(PermissionError):
            call(office, action, **({"document_id": key} if action == "read" else {}))


def test_existing_database_migration_preserves_chunks(office):
    store = office[3]
    with store.connect() as db:
        db.execute("ALTER TABLE chunks DROP COLUMN location")
        db.execute("INSERT INTO documents(path,size,mtime_ns,seen) VALUES('old.md',1,1,1)")
        db.execute("INSERT INTO chunks(document_id,revision,ordinal,text,tokens) VALUES(1,'r',0,'old','old')")
    store.initialize()
    store.initialize()
    with store.connect() as db:
        assert tuple(db.execute("SELECT text,location FROM chunks").fetchone()) == ("old", "{}")


def test_parser_cancellation_and_corrupt_office_are_visible(office):
    make_sheet(office[1] / "ok.xlsx")
    (office[1] / "bad.docx").write_bytes(b"broken")
    docs = ingest(office)
    assert next(doc for doc in docs if doc["path"] == "bad.docx")["state"] == "error"
    stop = threading.Event()
    stop.set()
    with pytest.raises(Interrupted):
        list(passages(office[1] / "ok.xlsx", office[4], stop, []))


def test_legacy_doc_missing_converter_reports_actionable_error(office, monkeypatch):
    monkeypatch.setenv("AGENTPARK_LIBREOFFICE", str(office[1] / "missing.exe"))
    (office[1] / "old.doc").write_bytes(b"legacy")
    docs = ingest(office)
    assert docs[0]["state"] == "error"
    assert "AGENTPARK_LIBREOFFICE" in docs[0]["error"]


def test_real_xls_read_and_group_aggregation(office):
    import shutil
    fixture = Path(__file__).parent / "fixtures/knowledge/legacy.xls"
    shutil.copyfile(fixture, office[1] / "legacy.xls")
    key = ingest(office)[0]["id"]
    preview = call(office, "read", document_id=key, sheet="Legacy", offset=1, limit=1)
    assert preview["rows"][0]["cells"][1]["value"] == 10
    assert "XLS" in preview["warnings"][1]
    result = call(office, "table", document_id=key, sheet="Legacy", start_row=2, group_by=["A"],
                  metrics=[dict(name="total", op="sum", column="B")])
    assert result["groups"][0]["metrics"]["total"]["value"] == "30.0"


def test_real_legacy_doc_conversion_and_original_unchanged(office):
    import shutil
    from src.knowledge.legacy_word import converter_path
    try:
        converter_path()
    except ValueError:
        pytest.skip("LibreOffice is not installed")
    source = office[1] / "legacy.doc"
    shutil.copyfile(Path(__file__).parent / "fixtures/knowledge/legacy.doc", source)
    before = source.read_bytes()
    doc = ingest(office)[0]
    assert doc["state"] == "embedding", doc["error"]
    result = call(office, "read", document_id=doc["id"])
    assert any("Hybrid search" in item["text"] for item in result["passages"])
    assert source.read_bytes() == before


def test_explicit_header_is_carried_into_later_chunks(office):
    _, folder, _, store, config = office
    book = Workbook()
    book.active.append(["Product", "Revenue"])
    for index in range(80):
        book.active.append([f"item-{index}", index])
    book.save(folder / "large.xlsx")
    config = config.model_copy(update={"spreadsheet_header_row": 1})
    warnings = []
    chunks = list(passages(folder / "large.xlsx", config, threading.Event(), warnings))
    assert len(chunks) < 80
    assert "Revenue" in chunks[-1].text
    assert chunks[-1].location["header_row"] == 1


def test_group_and_cell_limits_fail_without_partial_result(office, monkeypatch):
    book = Workbook()
    for index in range(201):
        book.active.append([f"group-{index}", index])
    book.save(office[1] / "groups.xlsx")
    key = ingest(office)[0]["id"]
    with pytest.raises(ValueError, match="分组超过"):
        call(office, "table", document_id=key, sheet="Sheet", group_by=["A"],
             metrics=[dict(name="rows", op="count_rows")])
    monkeypatch.setattr("src.knowledge.workbook_reader.MAX_CELLS", 10)
    with pytest.raises(ValueError, match="单元格"):
        call(office, "table", document_id=key, sheet="Sheet", metrics=[dict(name="rows", op="count_rows")])


def test_excel_merged_cells_and_error_values_are_not_filled_or_ignored(office):
    book = Workbook()
    book.active.append(["Region", "Amount"])
    book.active.append(["East", 3])
    book.active.append([None, "#DIV/0!"])
    book.active.merge_cells("A2:A3")
    book.save(office[1] / "merged.xlsx")
    doc = ingest(office)[0]
    assert "错误值" in doc["warning"]
    rows = call(office, "read", document_id=doc["id"], sheet="Sheet", offset=1)["rows"]
    assert rows[1]["cells"][0]["value"] is None
    with pytest.raises(ValueError, match="B3.*错误"):
        call(office, "table", document_id=doc["id"], sheet="Sheet", start_row=2,
             metrics=[dict(name="sum", op="sum", column="B")])


@pytest.mark.parametrize("arguments", [
    {"metrics": [dict(name="sum", op="sum")]},
    {"metrics": [dict(name="count", op="count_rows", column="A")]},
    {"metrics": [dict(name="sum", op="sum", column="a")]},
    {"metrics": [dict(name="sum", op="sum", column="XFE")]},
    {"metrics": [dict(name="n", op="count_rows")], "start_row": 5, "end_row": 2},
    {"metrics": [dict(name="n", op="count_rows")], "filters": [dict(column="A", op="eq")]},
    {"metrics": [dict(name="n", op="count_rows")], "expression": "__import__('os')"},
])
def test_table_protocol_rejects_invalid_or_executable_arguments(office, arguments):
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        call(office, "table", document_id=1, sheet="Sheet", **arguments)


def test_script_manifest_accepts_nested_table_contract_and_rejects_extra_fields():
    import jsonschema
    manifest = json.loads((Path(__file__).parents[1] / ".agents/skills/knowledge/skill.json").read_text(encoding="utf-8"))
    schema = next(item["argsSchema"] for item in manifest["scripts"] if item["id"] == "table")
    arguments = dict(library_id="x", document_id=1, sheet="Sheet", metrics=[dict(name="rows", op="count_rows")])
    jsonschema.validate(arguments, schema)
    arguments["metrics"][0]["code"] = "print(1)"
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(arguments, schema)


def test_large_xlsx_computes_entire_range_without_indexed_chunks(office):
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Data")
    sheet.append(["ID", "Amount"])
    for index in range(1, 20001):
        sheet.append([index, 1])
    book.save(office[1] / "large.xlsx")
    store = office[3]
    begin_scan(store)
    scan(store, office[1], threading.Event())
    key = store.documents()["items"][0]["id"]
    result = call(office, "table", document_id=key, sheet="Data", start_row=2,
                  metrics=[dict(name="total", op="sum", column="B"), dict(name="rows", op="count_rows")])
    assert result["complete"] and result["matched_rows"] == 20000
    assert result["groups"][0]["metrics"]["total"]["value"] == "20000"
    assert result["index_current"] is False
    assert result["end_row"] == 20001
    with store.connect() as db:
        assert db.execute("SELECT count(*) FROM chunks").fetchone()[0] == 0


def test_registered_path_cannot_escape_library_root(office):
    workspace, _, _, store, _ = office
    outside = workspace / "outside.txt"
    outside.write_text("not part of this library", encoding="utf-8")
    with store.connect() as db:
        db.execute("INSERT INTO documents(path,size,mtime_ns,seen) VALUES(?,?,?,1)",
                   ("../outside.txt", outside.stat().st_size, outside.stat().st_mtime_ns))
    with pytest.raises(ValueError, match="超出知识库"):
        call(office, "read", document_id=1)


def test_table_tool_runs_through_real_skill_subprocess(office, monkeypatch):
    import shutil
    from src.skills.script_manifest import load_skill_script_manifest, run_skill_script
    workspace, folder, library_id, _, _ = office
    root = Path(__file__).parents[1]
    shutil.copyfile(root / "tests/fixtures/knowledge/legacy.xls", folder / "legacy.xls")
    document_id = ingest(office)[0]["id"]
    skill_dir = workspace / "skills/knowledge"
    shutil.copytree(root / ".agents/skills/knowledge", skill_dir)
    monkeypatch.setenv("PYTHONPATH", str(root))
    definition = next(item for item in load_skill_script_manifest(str(skill_dir), skill_name="knowledge")
                      if item.id == "table")
    result = json.loads(run_skill_script(definition, dict(library_id=library_id, document_id=document_id,
                        sheet="Legacy", start_row=2, metrics=[dict(name="total", op="sum", column="B")])))
    assert result["status"] == "success", result
    body = json.loads(result["stdout"])
    assert body["groups"][0]["metrics"]["total"]["value"] == "37.0"
